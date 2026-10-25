"""Append identical V26 chapters while preserving existing report text and images."""
from pathlib import Path
import json,hashlib,shutil,datetime,zipfile,sys
from docx import Document
P=Path(__file__).resolve().parent;D=P/'documents';sys.path.insert(0,str(D))
import build_documents as b
from validation_decision_section import add_validation_decision
status=json.loads((D/'document_status.json').read_text(encoding='utf-8-sig'));newrows=[];regressions=[]
current='Current update on 25 October 2026: The evidence review is complete. Neither SoftSpace nor SoftSpace Flex has passed full software validation. Joint strength, complete product models and exact Sony mounting remain unresolved. The final Software validation decision page identifies what is established and what must still be done. No new simulation campaign or product pass is claimed.'
for name in ['Ear_Defender_Project.docx','Glasses_Friendly_Headphone_Cushions.docx']:
 f=D/name;oldhash=hashlib.sha256(f.read_bytes()).hexdigest();record=next(r for r in status['documents'] if Path(r['path']).name==name);assert oldhash==record['sha256'],'User changes detected; do not overwrite'
 d=Document(f);assert not any(p.text.startswith('Software validation decision') for p in d.paragraphs),'Decision already added'
 oldpars=[p.text for p in d.paragraphs if not p.text.startswith('Current update on')];oldtables=[[c.text for row in t.rows for c in row.cells] for t in d.tables]
 with zipfile.ZipFile(f) as z:oldmedia={hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')}
 backup=D/'history'/(f.stem+'_V26_before_decision_'+datetime.datetime.now().strftime('%Y%m%d_%H%M%S')+'.docx');shutil.copy2(f,backup)
 for p in d.paragraphs:
  if p.text.startswith('Current update on'):p.text=current
 add_validation_decision(d,P,b.new_page,b.para,b.table,b.picture)
 for section in d.sections:
  for hp in section.header.paragraphs:
   for run in hp.runs:
    if '24 October 2026' in run.text:run.text=run.text.replace('24 October 2026','25 October 2026')
 d.save(f);new=Document(f);newpars=[p.text for p in new.paragraphs if not p.text.startswith('Current update on')];assert newpars[:len(oldpars)]==oldpars
 assert [[c.text for row in t.rows for c in row.cells] for t in new.tables][:len(oldtables)]==oldtables
 with zipfile.ZipFile(f) as z:assert oldmedia.issubset({hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')});assert z.testzip() is None
 newrows.append(dict(path=str(f),sha256=hashlib.sha256(f.read_bytes()).hexdigest(),bytes=f.stat().st_size));regressions.append(dict(document=name,before_sha256=oldhash,all_old_paragraphs_preserved_except_current_summary=True,all_old_tables_preserved=True,all_old_images_preserved=True,backup=str(backup)))
status.update(updated='25 October 2026',design_status='Software validation decision',documents=newrows,native_Word_layout_verified=False)
(D/'document_status.json').write_text(json.dumps(status,indent=2));(D/'qa/decision_semantic_regression.json').write_text(json.dumps(regressions,indent=2))
# Keep future fresh rebuilds consistent with the directly patched reports.
builder=D/'build_documents.py';s=builder.read_text(encoding='utf-8-sig');lines=s.splitlines()
for i,line in enumerate(lines):
 if "para(d,'Current update on" in line:lines[i]=" para(d,"+repr(current)+")"
s='\n'.join(lines)+'\n'
for name in ['Ear_Defender_Project.docx','Glasses_Friendly_Headphone_Cushions.docx']:
 needle=" return save(d,'"+name+"')";assert needle in s;s=s.replace(needle," from validation_decision_section import add_validation_decision\n add_validation_decision(d,P,new_page,para,table,picture)\n"+needle)
s=s.replace("existing=Path(item['path'])","existing=O/Path(item['path']).name")
s=s.replace("hp.text='SOFTSPACE    |    24 October 2026'","hp.text='SOFTSPACE    |    25 October 2026'")
s=s.replace('def main_doc():',"REV='Evidence review complete with open product validation requirements'\n\ndef main_doc():",1);builder.write_text(s,encoding='utf8',newline='\n')
print(json.dumps(newrows,indent=2))
