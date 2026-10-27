from pathlib import Path
import json,hashlib,shutil,sys,zipfile
from docx import Document
P=Path(__file__).resolve().parent;D=P/'documents';C=P/'project-closure-2026-10-27';sys.path.insert(0,str(D));import build_documents as b
section="""from build_documents import link

def add_project_closure(d,P,new_page,para,table,picture):
 new_page(d,'Final saved CAD and project closure')
 para(d,'27 October 2026. Both projects are closed as research projects at the owner request. The archive preserves designs, assumptions, successes and failures. Completing this archive does not mean every software test passed or either product was physically validated. No further product development is planned.')
 picture(d,'V27_prototype_CAD','Latest saved STEP surfaces. SoftSpace has six full assemblies. Flex shows a material matched foam reference; the Sony mounting skirt is not modelled.',6.55)
 table(d,['Final saved work','Recorded status'],[['SoftSpace V27 assemblies','Six sizes with 42 solids each'],['Assembly adjustment and folding','189 sampled cases passed across six sizes'],['Wearing stop queries','Recorded blocking checks passed'],['Joint strength','Unresolved; no new accepted V26B equilibrium field'],['Native Fusion feature archives','Earlier V10 only; V27 is STEP']],[4.1,2.8])
 para(d,'V27 integrates the rounded relieved joint into full headset geometry. Finite rigid samples do not prove continuous clearance, manufactured fit, loaded retention or durability. A separate full to local correspondence check was not completed because its launch failed to import the project module. No result is claimed for that check.')
 new_page(d,'Final glasses cushion findings')
 picture(d,'V27_alignment_sensitivity','One declared comparison from the saved calculation. Negative values mean lower model pressure, not measured pain relief.',6.55)
 table(d,['Zone','Comparable cases','Pressure lower','Pressure higher'],[['18 mm','400','279','121'],['32 mm','399','399','0']],[1.5,1.8,1.8,1.8])
 para(d,'The 840 alignment scenarios include 799 comparable pairs and 41 without a paired solution in the supported material range. The wider zone is an experimental candidate. These assumed forces, arm dimensions and positions use typical sourced CF 42M and CF 40M loading curves in uncovered 12.7 mm foam. They are not people, success probabilities, a cover model or proof of comfort.')
 para(d,'Both new Flex references contain three valid nonoverlapping solids and a continuous nominal front cover. Exact Sony mounting remains unverified. Reusing the skirt of a fitting donor pad was proposed but not built or tested. There are no measured sound leakage, bass, isolation, ANC or durability results.')
 new_page(d,'Final conclusions and research context')
 para(d,'SoftSpace began as a passive mechanical headset study. The later ambition was premium quality ANC below 5000 rupees. No ANC electronics, controller or tuned microphone and speaker system was developed. Mechanical CAD and passive screens cannot satisfy that electronic target. Neither performance nor selling cost was established.')
 para(d,'External research describes personalised noise control for autistic children. A 2025 laboratory paper used tunable hybrid ANC. A related 2026 paper reports that 18 of 24 parents observed improved reactions after three weeks. These are not SoftSpace results and must not be counted as independent proof that most autistic people prefer ANC.')
 link(d,'Tunable ANC research in Scientific Reports 2025','https://www.nature.com/articles/s41598-025-89609-6')
 link(d,'Related 2026 headset study and parent report abstract','https://research.polyu.edu.hk/en/publications/auditory-perception-based-tuneable-noise-control-approach-for-imp/')
 para(d,'For glasses, a 2019 fixture study found that wide spectacle arms disrupted earmuff isolation more than thin arms. This supports the seal problem, but does not establish exact consumer headphone losses or prove our pad works. Turtle Beach already has adjustable glasses relief; it is not merely ordinary memory foam. No absence of Indian competition or patent clearance is claimed.')
 link(d,'Peer reviewed earmuff and eyewear study 2019','https://pmc.ncbi.nlm.nih.gov/articles/PMC6539177/')
 link(d,'Turtle Beach adjustable glasses relief instructions','https://support.turtlebeach.com/articles/en_US/Knowledge/115001384573-Elite-Pro-PC-Edition-How-to-Remove-Ear-Cushions-Adjust-ProSpecs-Glasses-Relief')
 para(d,'The retained joint checkpoint has relative force residual about 0.00000754, above the strict final target 0.0000001. Small independent force and moment sums do not override this failed gate or establish stress accuracy. Later solver failures and a diagnostic cleanup error are retained in the execution records. No result is promoted to a strength pass.')
 para(d,'Final outcome: completed research archive with useful CAD and conditional comparisons, unresolved validation and no physical prototype tests. Startup development is discontinued. Previous next steps and purchase plans are historical records, not active instructions. The owner can explicitly reopen either project later.')
"""
(D/'project_closure_section.py').write_text(section);exec(compile(section,'project_closure_section.py','exec'))
status=json.loads((D/'document_status.json').read_text());records=[];reg=[]
current='Project closed on 27 October 2026: Both ideas are retained as research projects. Latest V27 CAD and calculations are archived, including failures. No physical testing or ANC system was completed. Read the final closure pages; earlier development plans are historical.'
for name in ['Ear_Defender_Project.docx','Glasses_Friendly_Headphone_Cushions.docx']:
 f=D/name;oldhash=hashlib.sha256(f.read_bytes()).hexdigest();assert oldhash==next(r['sha256'] for r in status['documents'] if Path(r['path']).name==name),'User edits detected'
 d=Document(f);assert not any(p.text=='Final saved CAD and project closure' for p in d.paragraphs);oldpars=[p.text for p in d.paragraphs if not p.text.startswith('Current update on')];tables=[[c.text for row in t.rows for c in row.cells] for t in d.tables]
 with zipfile.ZipFile(f) as z:media={hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')}
 backup=D/'history'/(f.stem+'_before_closure_20261027.docx');assert not backup.exists();shutil.copy2(f,backup)
 for p in d.paragraphs:
  if p.text.startswith('Current update on'):p.text=current
 add_project_closure(d,P,b.new_page,b.para,b.table,b.picture)
 for sec in d.sections:
  for hp in sec.header.paragraphs:
   for run in hp.runs:run.text=run.text.replace('26 October 2026','27 October 2026')
 d.save(f);n=Document(f);assert [p.text for p in n.paragraphs if not p.text.startswith('Project closed on')][:len(oldpars)]==oldpars;assert [[c.text for row in t.rows for c in row.cells] for t in n.tables][:len(tables)]==tables
 with zipfile.ZipFile(f) as z:assert media.issubset({hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')});assert z.testzip() is None
 records.append(dict(path=str(f),sha256=hashlib.sha256(f.read_bytes()).hexdigest(),bytes=f.stat().st_size));reg.append(dict(document=name,before_sha256=oldhash,prior_text_preserved_except_current_summary=True,prior_tables_and_media_preserved=True,backup=str(backup)))
status.update(updated='27 October 2026',design_status='Final saved CAD and project closure',documents=records,project_status='CLOSED_RESEARCH_PROJECT',native_Word_layout_verified=False,closure_date='2026-10-27');status['quality_check']['native_Word_layout_verified']=False;(D/'document_status.json').write_text(json.dumps(status,indent=2));(D/'qa/closure_semantic_regression.json').write_text(json.dumps(reg,indent=2))
f=D/'build_documents.py';s=f.read_text(encoding='utf-8-sig')
for name in ['Ear_Defender_Project.docx','Glasses_Friendly_Headphone_Cushions.docx']:
 needle=" return save(d,'"+name+"')";assert needle in s;s=s.replace(needle," from project_closure_section import add_project_closure\n add_project_closure(d,P,new_page,para,table,picture)\n".replace('\n','\n')+needle)
# Explicit plain replacement avoids replacing historical chapter prose.
lines=s.splitlines();lines=[(' para(d,'+repr(current)+')') if "para(d,'Current update on" in line else line for line in lines];s='\n'.join(lines)+'\n';s=s.replace("hp.text='SOFTSPACE    |    26 October 2026'","hp.text='SOFTSPACE    |    27 October 2026'");f.write_text(s,encoding='utf-8')
f=D/'check_native_word_layout_block.py';(D/'check_native_word_layout_closure.py').write_text(f.read_text().replace('native_block_','native_closure_').replace("date='2026-10-26'","date='2026-10-27'"))
f=D/'verify_documents.py';f.write_text(f.read_text().replace('native_block_layout_checks.json','native_closure_layout_checks.json'))
f=D/'export_native_word.ps1';f.write_text(f.read_text().replace("date='2026-10-26'","date='2026-10-27'"))
(C/'execution_note.md').write_text('The near-equilibrium field is explicitly diagnostic with converged=false. The original diagnostic saved its JSON and NPZ before a cleanup NameError. A later cleanup-only edit was not rerun and is not a new successful solve. No new simulations were run for closure. Full-to-local geometry correspondence was not checked because the launch failed to import the project module.\n')
print('Both final reports saved')
