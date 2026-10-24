"""Independent approximate layout check. This is NOT Word/LibreOffice rendering."""
from pathlib import Path
from io import BytesIO
import json,html,zipfile,hashlib
from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph as DP
from docx.table import Table as DT
from reportlab.platypus import SimpleDocTemplate,Paragraph,Image,Table,TableStyle,PageBreak,Spacer,KeepTogether
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor
import pypdfium2 as pdfium
O=Path(__file__).parent;Q=O/'qa';Q.mkdir(exist_ok=True)
for name,file in [('Calibri','calibri.ttf'),('CalibriBold','calibrib.ttf'),('CalibriItalic','calibrii.ttf'),('CalibriBI','calibriz.ttf')]:
 pdfmetrics.registerFont(TTFont(name,'C:/Windows/Fonts/'+file))
pdfmetrics.registerFontFamily('Calibri',normal='Calibri',bold='CalibriBold',italic='CalibriItalic',boldItalic='CalibriBI')
results=[]
for f in [O/'Ear_Defender_Project.docx',O/'Glasses_Friendly_Headphone_Cushions.docx']:
 d=Document(f);flows=[];text=[];pics=0;tables=0
 for node in d.element.body:
  if node.tag==qn('w:p'):
   p=DP(node,d)
   if node.xpath('.//w:br[@w:type="page"]'):flows.append(PageBreak());continue
   blips=node.xpath('.//a:blip')
   if blips:
    extent=node.xpath('.//wp:extent')[0];image=d.part.related_parts[blips[0].get(qn('r:embed'))].blob
    flows.append(Image(BytesIO(image),width=int(extent.get('cx'))/12700,height=int(extent.get('cy'))/12700));pics+=1;continue
   t=''.join(node.xpath('.//w:t/text()'));text.append(t)
   if not t:flows.append(Spacer(1,6));continue
   
   if p.paragraph_format.page_break_before:flows.append(PageBreak())
   st=p.style.name;size=28 if st=='Title' else (19 if st=='Heading 1' else (14 if st=='Subtitle' else (9 if st=='Caption' else 11)))
   if p.runs and p.runs[0].font.size:size=p.runs[0].font.size.pt
   ps=ParagraphStyle('p',fontName='CalibriBold' if st in ['Title','Heading 1'] else ('CalibriItalic' if st=='Caption' else 'Calibri'),fontSize=size,leading=size*1.12,spaceAfter=7,keepWithNext=st.startswith('Heading'))
   flows.append(Paragraph(html.escape(t),ps))
  elif node.tag==qn('w:tbl'):
   tb=DT(node,d);rows=[];tables+=1
   for ri,row in enumerate(tb.rows):
    cells=[]
    for c in row.cells:
     txt='\n'.join(p.text for p in c.paragraphs);text.append(txt)
     cells.append(Paragraph(html.escape(txt).replace('\n','<br/>'),ParagraphStyle('cell',fontName='CalibriBold' if ri==0 else 'Calibri',fontSize=10,leading=10.8,textColor=HexColor('#ffffff' if ri==0 else '#000000'))))
    rows.append(cells)
   widths=[c.width.pt for c in tb.columns];t=Table(rows,colWidths=widths,repeatRows=1,hAlign='CENTER')
   t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),HexColor('#203e4b')),('ROWBACKGROUNDS',(0,1),(-1,-1),[HexColor('#f0f5f5'),HexColor('#ffffff')]),('GRID',(0,0),(-1,-1),.5,HexColor('#d9d9d9')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('LEFTPADDING',(0,0),(-1,-1),4.5),('RIGHTPADDING',(0,0),(-1,-1),4.5),('TOPPADDING',(0,0),(-1,-1),4.5),('BOTTOMPADDING',(0,0),(-1,-1),4.5)]));flows.append(t)
 out=Q/f.stem;out.mkdir(exist_ok=True);pdf=out/'approximate_layout.pdf'
 def footer(canvas,doc):
  canvas.setFont('Calibri',8);canvas.drawString(57.6,766,'SOFTSPACE  |  '+json.loads((O/'document_status.json').read_text())['updated']);canvas.drawRightString(554.4,28,'Approximate internal layout check  |  '+str(doc.page))
 SimpleDocTemplate(str(pdf),pagesize=(612,792),leftMargin=57.6,rightMargin=57.6,topMargin=46.8,bottomMargin=46.8).build(flows,onFirstPage=footer,onLaterPages=footer)
 pd=pdfium.PdfDocument(str(pdf));pages=[]
 for i in range(len(pd)):
  target=out/('page-'+str(i+1)+'.png');pd[i].render(scale=1.5).to_pil().save(target);pages.append(str(target))
 with zipfile.ZipFile(f) as z:assert z.testzip() is None
 result={'file':str(f),'sha256':hashlib.sha256(f.read_bytes()).hexdigest(),'approximate_pages':len(pd),'images':pics,'tables':tables,'zip_integrity':'pass','native_page_render':'unverified: bundled LibreOffice missing; Word control returned unavailable controls and activation failure','plain_text':'\n'.join(text)}
 (out/'content.txt').write_text(result.pop('plain_text'),encoding='utf-8');results.append(result);print({k:v for k,v in result.items() if k!='file'},flush=True)
(Q/'checks.json').write_text(json.dumps(results,indent=2))
