"""Native Word pagination fixes shared by both living reports."""
from docx.shared import Pt,Inches
from docx.oxml.ns import qn

def apply_layout(d):
 for border in list(d.styles['Title'].element.xpath('.//w:pBdr')):border.getparent().remove(border)
 d.styles['Title'].font.underline=False
 d.styles['Normal'].font.size=Pt(10.5)
 d.styles['Normal'].paragraph_format.line_spacing=1.05
 d.styles['Normal'].paragraph_format.space_after=Pt(4)
 d.styles['Heading 1'].paragraph_format.space_before=Pt(0)
 d.styles['Heading 1'].paragraph_format.space_after=Pt(8)
 chapter=''
 for p in list(d.paragraphs):
  if p.style.name=='Heading 1':chapter=p.text
  if any(chapter.startswith(t) for t in ['V13 Assembly failure','V15 Shared packing','V16 Joint clearance','V18 Geometry and rejected trials','V18 Retention result']):
   scale=.82 if chapter.startswith('V18 Geometry') else .85 if chapter.startswith('V16 Joint clearance') else .9
   for extent in p._p.xpath('.//wp:extent | .//a:xfrm/a:ext'):
    for key in ['cx','cy']:extent.set(key,str(round(int(extent.get(key))*scale)))
   if chapter.startswith(('V18 Geometry','V16 Joint clearance')):
    if p.style.name=='Caption':p.paragraph_format.space_after=Pt(4)
    if not p.text.strip() and not p._p.xpath('.//w:drawing') and p._p.getparent() is not None:
     p._p.getparent().remove(p._p);continue
  if p._p.xpath('.//w:br[@w:type="page"]') and not p.text.strip():
   nxt=p._p.getnext()
   if nxt is not None and nxt.tag==qn('w:p'):
    from docx.text.paragraph import Paragraph
    Paragraph(nxt,d).paragraph_format.page_break_before=True
    p._p.getparent().remove(p._p)
  if p._p.xpath('.//w:hyperlink'):
   p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1.0
   for n in p._p.xpath('.//w:r'):
    from docx.text.run import Run
    Run(n,p).font.size=Pt(9)
 for t in d.tables:
  widths=[c.width.inches for c in t.columns];scale=min(1,6.9/sum(widths));widths=[v*scale for v in widths]
  for c,v in zip(t.columns,widths):c.width=Inches(v)
  for row in t.rows:
   for c,v in zip(row.cells,widths):
    c.width=Inches(v)
    for p in c.paragraphs:
     p.paragraph_format.line_spacing=1.03;p.paragraph_format.space_after=Pt(2)
