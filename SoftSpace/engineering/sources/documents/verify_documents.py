from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
import json,hashlib,zipfile,re
O=Path(__file__).parent;status=json.loads((O/'document_status.json').read_text());checks=[]
for item in status['documents']:
 f=Path(item['path']);d=Document(f);assert hashlib.sha256(f.read_bytes()).hexdigest()==item['sha256']
 with zipfile.ZipFile(f) as z:assert z.testzip() is None
 text='\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(c.text for t in d.tables for r in t.rows for c in r.cells)
 links=[str(r.target_ref) for r in d.part.rels.values() if r.reltype.endswith('/hyperlink')]
 assert len(links)>=6 and status['design_status'] in text
 assert all(im._inline.docPr.get('descr') for im in d.inline_shapes)
 assert all(t.rows[0]._tr.xpath('./w:trPr/w:tblHeader') for t in d.tables)
 assert not d.styles['Title'].element.xpath('.//w:pBdr')
 assert all(not p._p.xpath('./w:pPr/w:pBdr') for p in d.paragraphs if p.style.name=='Title')
 if f.stem=='Ear_Defender_Project':
  last=10 if (O.parent/'engineering-v10-results'/'cad_local_report.json').exists() else 9
  assert all(any(p.text.startswith('V'+str(n)+' ') and p.style.name=='Heading 1' for p in d.paragraphs) for n in range(1,last+1))
 checks.append({'document':f.name,'paragraphs':len(d.paragraphs),'tables':len(d.tables),'embedded_images':len(d.inline_shapes),'source_links':len(links),'structure_and_hashes':'pass'})
qa=json.loads((O/'qa'/'checks.json').read_text()) if (O/'qa'/'checks.json').exists() else []
matched=all(any(q.get('sha256')==item['sha256'] for q in qa) for item in status['documents'])
status['quality_check']={'structural_checks':'pass','approximate_layout_page_counts':[q['approximate_pages'] for q in qa] if matched else [],'approximate_previews_match_current_documents':matched,'native_Word_layout_verified':False,'reason':'Bundled LibreOffice was unavailable in the recorded renderer check. Native desktop document control is not available for this batch. Approximate previews were used; native Word pagination remains unverified.','checks':checks}
native=O/'qa/native_word_status.json';layout=O/'qa/native_v22_layout_checks.json'
if native.exists() and layout.exists():
 n=json.loads(native.read_text(encoding='utf-8-sig'));l=json.loads(layout.read_text());native_match=all(any(r['source_sha256']==item['sha256'] for r in n['documents']) for item in status['documents']);layout_match=all(any(r['source_sha256']==item['sha256'] for r in l['documents']) for item in status['documents'])
 if native_match and layout_match:
  status['quality_check'].update(native_Word_export_verified=True,native_page_counts=[r['native_pages'] for r in n['documents']],native_blank_or_out_of_page_text_pages=sum(len(r['blank_body_pages'])+r['out_of_page_text_rectangles'] for r in l['documents']),native_Word_layout_verified=False,reason='Native Word read-only PDF exports match the current DOCX hashes. All-page render, page bounds and blank-page checks completed. Visual review is recorded separately.')
(O/'document_status.json').write_text(json.dumps(status,indent=2));print(json.dumps(checks,indent=2))

